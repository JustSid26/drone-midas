"""Builds the Unit 1 and 2 project presentation, styled to match the team's
existing autonomousShuttle deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

W, H = 20.0, 11.25
INK   = RGBColor(0x40, 0x3E, 0x3E)
MUTED = RGBColor(0x8A, 0x86, 0x84)
BODY  = RGBColor(0x55, 0x52, 0x52)
RULE  = RGBColor(0xD5, 0xD2, 0xCE)
CARD  = RGBColor(0xFD, 0xFD, 0xFA)
BG    = RGBColor(0xEF, 0xEE, 0xEC)
ONDK  = RGBColor(0xEF, 0xEE, 0xEC)
ONDKB = RGBColor(0xC9, 0xC6, 0xC3)
F, FS, FL = "Inter", "Inter SemiBold", "Inter Light"

M, CW = 0.9, 18.2
TOP = 2.25
BOT = 9.95

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
_n = [0]


def _bg(slide):
    """Paint the slide background. The element must sit first inside cSld, not
    as a sibling of it, or the file is schema invalid and Keynote refuses it."""
    from pptx.oxml.ns import qn
    from lxml import etree
    cSld = slide._element.find(qn('p:cSld'))
    bg = etree.Element(qn('p:bg'))
    pr = etree.SubElement(bg, qn('p:bgPr'))
    fill = etree.SubElement(pr, qn('a:solidFill'))
    clr = etree.SubElement(fill, qn('a:srgbClr'))
    clr.set('val', 'EFEEEC')
    etree.SubElement(pr, qn('a:effectLst'))
    cSld.insert(0, bg)


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.01):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                               Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw * 72)
    s.text_frame.text = ""
    return s


def tbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    b = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def line(tf, text, size, color, font=F, first=False, space_before=0, space_after=4,
         align=PP_ALIGN.LEFT, spacing=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = spacing
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.color.rgb = color
    return p


def slide(eyebrow, title):
    s = prs.slides.add_slide(BLANK)
    _bg(s)
    tf = tbox(s, M, 0.63, 18.3, 1.35)
    line(tf, eyebrow.upper(), 15, MUTED, FL, first=True, space_after=6)
    line(tf, title, 43.5, INK, F, space_after=0)
    rect(s, M, 1.92, CW, 0.02, fill=RULE)
    _n[0] += 1
    num = tbox(s, W - 1.6, H - 0.85, 0.7, 0.35)
    line(num, str(_n[0]), 13, MUTED, FL, first=True, align=PP_ALIGN.RIGHT)
    return s


def card(s, x, y, w, h, heading, lines, dark=False, hsize=21, bsize=16.5, gap=7):
    rect(s, x, y, w, h, fill=(INK if dark else CARD),
         line=(None if dark else RULE))
    tf = tbox(s, x + 0.42, y + 0.36, w - 0.84, h - 0.72)
    line(tf, heading, hsize, ONDK if dark else INK, FS, first=True, space_after=gap)
    for t in lines:
        line(tf, t, bsize, ONDKB if dark else BODY, F, space_after=gap, spacing=1.12)
    return tf


def arrows(s, x, y, w, h, items, size=19, gap=11):
    tf = tbox(s, x, y, w, h)
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.14
        r = p.add_run(); r.text = "—>   "
        r.font.name = F; r.font.size = Pt(size - 3); r.font.color.rgb = MUTED
        r2 = p.add_run(); r2.text = t
        r2.font.name = F; r2.font.size = Pt(size); r2.font.color.rgb = BODY
    return tf


def table(s, x, y, w, h, headers, rows, col_w=None, hsize=14, bsize=13):
    shape = s.shapes.add_table(len(rows) + 1, len(headers),
                               Inches(x), Inches(y), Inches(w), Inches(h))
    t = shape.table
    t.first_row = True
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    for j, htxt in enumerate(headers):
        c = t.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = INK
        c.margin_left = c.margin_right = Inches(0.12)
        c.margin_top = c.margin_bottom = Inches(0.07)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.space_after = Pt(0)
        r = p.add_run(); r.text = htxt
        r.font.name = FS; r.font.size = Pt(hsize); r.font.color.rgb = ONDK
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = CARD if i % 2 else BG
            c.margin_left = c.margin_right = Inches(0.12)
            c.margin_top = c.margin_bottom = Inches(0.06)
            c.vertical_anchor = MSO_ANCHOR.TOP
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.space_after = Pt(0); p.line_spacing = 1.08
            r = p.add_run(); r.text = str(val)
            r.font.name = F; r.font.size = Pt(bsize); r.font.color.rgb = BODY
    return t


def picture(s, path, x, y, w):
    return s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))

# =============================================================== 1. TITLE
s = prs.slides.add_slide(BLANK)
_bg(s)
rect(s, 1.65, 2.43, 2.25, 0.04, fill=INK)
tf = tbox(s, 1.65, 3.55, 16.8, 3.2)
line(tf, "Autonomous Drone Obstacle Avoidance", 62, INK, FS, first=True,
     space_after=2, spacing=1.02)
line(tf, "Using Monocular Depth Estimation", 62, INK, FS, space_after=10, spacing=1.02)
line(tf, "Fixed Wing VTOL  ·  Robotics and Computer Vision", 34, MUTED, FL, spacing=1.05)
rect(s, 1.65, 8.02, 4.5, 0.02, fill=RULE)
tf = tbox(s, 1.65, 8.45, 16.5, 1.4)
line(tf, "Siddharth Lama     Manit Aravind     Sahil Asopa     Riyaan Sheth",
     24, INK, FS, first=True, space_after=8)
line(tf, "BTECH CAP 501  Project Life Cycle Management  ·  Unit 1 and 2", 17, MUTED, FL)
_n[0] += 1
num = tbox(s, W - 1.6, H - 0.85, 0.7, 0.35)
line(num, "1", 13, MUTED, FL, first=True, align=PP_ALIGN.RIGHT)

# =============================================================== 2. INTRODUCTION
s = slide("Domain  ·  Robotics and Computer Vision", "Project Introduction")
arrows(s, M, TOP, 10.6, 6.2, [
    "Small drones now inspect towers, survey sites, monitor crops and carry parcels. "
    "Almost all of that work happens close to structures and close to people.",
    "To avoid obstacles you must know distance. Laser scanners measure it well but are "
    "heavy, costly and power hungry. Stereo pairs need calibration that survives vibration.",
    "On a small airframe the binding currency is mass. Every gram of sensor is a gram not "
    "spent on battery, and battery is flight time.",
    "One ordinary camera weighs a few grams and costs very little. It simply cannot "
    "measure distance, because projection destroys depth.",
])
card(s, 12.22, TOP, 6.9, 3.5, "The problem area",
     ["A single photograph has no distance in it. A small object near the lens and a "
      "large one far away project to identical pixels.",
      "No amount of geometry recovers the difference from one image."])
card(s, 12.22, 6.07, 6.9, 3.5, "What changed",
     ["Photographs of the real world are highly structured. Distant surfaces look smaller, "
      "lose contrast and are hidden by nearer ones.",
      "MiDaS learned those cues from a mixture of public datasets and generalises to "
      "scenes it has never seen."], dark=True)

# =============================================================== 3. PROBLEM AND SCOPE
s = slide("Why it matters  ·  what is in and out", "Problem Statement and Scope")
card(s, M, TOP, 10.6, 2.55, "The problem, stated precisely",
     ["A small aircraft carrying only one forward facing colour camera must reach a "
      "sequence of positions in a cluttered space without striking anything, using no "
      "distance sensor, no prior map and no external positioning."])
arrows(s, M, 5.15, 10.6, 4.6, [
    "Depth is predicted, not measured. A prediction can be confidently wrong and nothing "
    "in the output says when.",
    "The prediction has no units. It says which things are nearer, not that the wall is "
    "three metres away.",
    "The camera sees a wedge. A one metre wingspan sweeps a full circle. Anything that "
    "leaves the frame is invisible while the wingtip is still moving.",
], size=18)
card(s, 12.22, TOP, 6.9, 3.5, "In scope",
     ["Simulation as the graded deliverable, with perfect depth available for comparison.",
      "Pretrained MiDaS used as published. No training.",
      "A transparent reactive policy with a non learned safety stop underneath.",
      "Seeded courses so every depth source is judged on identical layouts."], bsize=15.5)
card(s, 12.22, 6.07, 6.9, 3.5, "Out of scope",
     ["Mapping and global path planning. The planner only ever reacts to the current frame.",
      "Training a depth network from scratch.",
      "Moving obstacles.",
      "Outdoor flight, until the regulatory position is confirmed in writing."],
     dark=True, bsize=15.5)

# =============================================================== 4. SOLUTION
s = slide("One learned stage, everything else readable", "Proposed AI and ML Solution")
picture(s, "figures/fig2_architecture.png", M + 0.15, TOP + 0.05, 5.2)
card(s, 6.6, TOP, 6.1, 3.5, "Uncertainty and experimentation",
     ["The network returns a plausible estimate, not a guaranteed value. Two similar "
      "scenes can give different depth.",
      "So the loop is built to tolerate it. The chosen heading is smoothed over time "
      "rather than taken from the newest frame."], bsize=15.5)
card(s, 6.6, 6.07, 6.1, 3.5, "Model evaluation",
     ["Depth quality on real photographs, using NYU and KITTI, after aligning both scale "
      "and offset by least squares.",
      "Flight outcomes across many held out seeded courses, reported as distributions "
      "with a confidence interval."], bsize=15.5)
card(s, 13.0, TOP, 6.1, 3.5, "Monitoring after deployment",
     ["Scene statistics such as brightness and near field pixel fraction are tracked "
      "against a reference distribution.",
      "Drift lowers the speed and raises the stop distance. Less confidence must always "
      "produce more caution."], dark=True, bsize=15.5)
card(s, 13.0, 6.07, 6.1, 3.5, "Retraining when conditions change",
     ["Deliberately a second line response, never the first.",
      "Light adaptation to new scenes must pass the simulation quality gate before it is "
      "promoted. Slowing down is safer than a fresh model nobody has tested."],
     dark=True, bsize=15.5)

# =============================================================== 5. LIFE CYCLE
s = slide("Problem to monitoring, with the loops that matter", "Life Cycle and ML Pipeline")
picture(s, "figures/fig1_lifecycle.png", M + 0.1, TOP - 0.05, 5.6)
card(s, 7.0, TOP, 5.7, 3.5, "The two feedback loops",
     ["Evaluation back to data and model. Poor ablation results send work to perception, "
      "not to the steering logic.",
      "Monitoring back to scoping. If scenes stop resembling development, the problem has "
      "moved even though nothing broke."], bsize=15.5)
card(s, 7.0, 6.07, 5.7, 3.5, "We have already used loop one",
     ["An early evaluation had the aircraft barely moving.",
      "The lineage record traced it to feature creation. Per column distance included the "
      "floor, so every direction read the same and the policy could not tell open space "
      "from an obstacle."], bsize=15.5)
card(s, 13.0, TOP, 6.1, 3.5, "Automated pipeline stages",
     ["Course generation from seeds, and frame export.",
      "Depth benchmarking and per seed flight scoring.",
      "Policy parameter search and drift comparison.",
      "The quality gate before a change is accepted."], bsize=15)
card(s, 13.0, 6.07, 6.1, 3.5, "Human judgement and accountability",
     ["What density is a fair test, not a flattering one.",
      "Which seeds are held back, fixed before results.",
      "What accuracy justifies flying at all.",
      "That it is fit to fly near people."],
     dark=True, bsize=15)

# =============================================================== 6. FRAMEWORKS
s = slide("CRISP-DM  ·  TDSP  ·  Google ML", "Comparing Life Cycle Frameworks")
table(s, M, TOP, 12.0, 4.6,
      ["Aspect", "CRISP-DM", "Microsoft TDSP", "Google ML Life Cycle"],
      [["Primary emphasis", "Understanding the problem through iteration",
        "Roles, deliverables and repeatable team process",
        "Production serving and continuous operation"],
       ["Data work", "Central, with its own phase and explicit rework",
        "Structured acquisition with a documented manifest",
        "Automated validation and transformation"],
       ["Operation", "Ends at deployment", "Ends at customer acceptance",
        "Assumes a running system with retraining"],
       ["Automation and pipeline", "Little. The cycle is human driven",
        "Moderate, through templates and reports",
        "Strong. The pipeline is the unit of work"],
       ["Team collaboration", "Light, no prescribed roles",
        "Explicit named owner per deliverable",
        "Assumes a platform team already exists"]],
      col_w=[2.2, 3.2, 3.2, 3.4], hsize=14, bsize=12.5)
card(s, 13.4, TOP, 5.7, 3.5, "Best fit for this project",
     ["CRISP-DM. Our work is investigative. We do not know in advance how much "
      "performance is lost when measured depth is replaced by predicted depth.",
      "It expects evaluation to send work back, which is exactly what happened to us."],
     bsize=15.5)
card(s, 13.4, 6.07, 5.7, 3.5, "Framework selection",
     ["CRISP-DM sets the shape of the work. Google pipeline discipline applies to anything "
      "that runs more than once.",
      "TDSP contributes only its named owner per deliverable. Its customer acceptance stage "
      "assumes a client we do not have."], dark=True, bsize=15.5)

# =============================================================== 7. TECHNICAL FEASIBILITY
s = slide("Tested, not asserted", "Technical Feasibility")
table(s, M, TOP, 12.0, 5.6,
      ["Aspect", "Project analysis", "Risk", "Mitigation"],
      [["Data sufficiency", "No training data needed. The simulator gives exact depth free",
        "Simulator imagery is cleaner than photographs",
        "Report benchmark accuracy on real images alongside"],
       ["Model capability", "MiDaS-small, 21 M parameters, built for CPU",
        "Output has no units, unknown scale and offset",
        "Fit both by least squares. Median scaling alone is wrong"],
       ["Latency and hardware", "Our loop sustains over 200 cycles per second on a laptop",
        "Streaming video over Wi-Fi can eat the whole 100 ms budget",
        "Measure the round trip. Benchmark the model on the Pi first"],
       ["Sensor coverage", "Camera sees a wedge. The wing is 1.00 m across",
        "Obstacles leaving the frame are hit by the wingtip",
        "Wide lens plus a short term decaying memory. Both tested"],
       ["Airframe control", "Four motors with two tilt servos, so a tilting rotor",
        "Tilt mechanisms and transition are the hardest VTOL case",
        "Keep all avoidance testing in hover. Transition is separate"]],
      col_w=[1.9, 3.5, 3.2, 3.4], hsize=14, bsize=12.5)
picture(s, "figures/fig5_results.png", 13.5, TOP + 0.1, 5.5)
card(s, 13.4, 6.5, 5.7, 3.1, "What the measurement says",
     ["Every figure uses perfect depth, so it is the ceiling. Estimated depth cannot beat it.",
      "Tuning on 12 courses gave 92 per cent. The same settings gave 66 per cent on 50 held "
      "out courses. Quote the second number."], dark=True, bsize=15)

# ========================================= 8. ECONOMIC, OPERATIONAL, ETHICAL
s = slide("Cost, operation and harm", "Economic, Operational and Ethical Feasibility")
table(s, M, TOP, 9.0, 3.3,
      ["Aspect", "Project analysis", "Risk", "Mitigation"],
      [["Economic", "Simulation costs nothing. Airframe 55,000 to 90,000 rupees",
        "Buying motors before the mass budget exists",
        "Two stage purchase. Propulsion last"],
       ["Economic", "No labelling needed. The simulator supplies exact depth",
        "Development effort is the real cost, paid in time",
        "Four owners, one per workstream"],
       ["Operational", "Team operates it. A named pilot holds a manual override",
        "An alert nobody reads is not monitoring",
        "Named owner for drift review"],
       ["Operational", "Model never retrained, so no retraining pipeline",
        "Work becomes unusable once the team disperses",
        "Seeded courses, pinned versions, one hour to reproduce"]],
      col_w=[1.5, 3.0, 2.3, 2.2], hsize=13.5, bsize=12)
table(s, M, 6.05, 9.0, 3.5,
      ["Ethical issue", "Affected group", "Potential harm", "Mitigation"],
      [["Privacy", "Anyone within view who never agreed",
        "A camera above head height records private ground",
        "No frame written to disk in flight. Nothing leaves the operator"],
       ["Bias and fairness", "People near textureless or thin structures",
        "The model is worse there, moving risk onto them",
        "Failure modes published, stop distance set for the worst case"],
       ["Misuse", "Bystanders and observed individuals",
        "Autonomy lowers the barrier to persistent observation",
        "No person detection, tracking or recording exists. Deliberate"],
       ["Oversight", "Bystanders and the operator",
        "Nobody accountable if the system is treated as responsible",
        "Named pilot in command. The stop layer cannot be overridden"]],
      col_w=[1.5, 2.2, 2.7, 2.6], hsize=13.5, bsize=12)
card(s, 10.3, TOP, 8.8, 3.3, "Economic conclusion",
     ["The graded deliverable is unambiguously feasible. It costs nothing beyond hardware "
      "the team owns and time the semester already allocates.",
      "The airframe is the only real expenditure, and propulsion is deferred until all up "
      "mass and thrust to weight are known."], bsize=16)
card(s, 10.3, 6.05, 8.8, 3.5, "Ethical conclusion. It may proceed, with conditions",
     ["Graded work stays in simulation. Flight begins indoors, enclosed, over soft flooring, "
      "guards fitted, override held.",
      "No transition until hover is stable and repeatable.",
      "Published results always quote estimated depth on held out courses. Overstating "
      "reliability for a system that can strike a person is a safety failure, not a "
      "presentational one."], dark=True, bsize=16)

# =============================================================== 9. REQUIREMENTS
s = slide("Written so each one can pass or fail", "Requirements Engineering")
card(s, M, TOP, 8.9, 3.5, "Functional",
     ["FR3  A safety layer refuses forward motion whenever the nearest surface at flying "
      "height is inside 1.2 m, overriding the policy.",
      "FR4  The aircraft flies an ordered sequence of operator specified stations with no "
      "further human input.",
      "FR6  A given integer seed always produces the identical obstacle layout."], bsize=16)
card(s, M, 6.07, 8.9, 3.5, "Non functional",
     ["NFR3  Camera to command latency at most 100 ms at 2 m/s. Derived below, not chosen.",
      "NFR4  Zero contacts across 20 runs with the steering policy disabled and only the "
      "stop rule active.",
      "NFR5  At least 0.80 of pixels within a factor of 1.25 of true depth on NYU, after "
      "aligning scale and offset."], bsize=16)
card(s, 10.2, TOP, 8.9, 3.5, "Data",
     ["At least 5,000 colour and true depth pairs from seeded courses, plus the NYU and "
      "KITTI evaluation splits.",
      "Camera intrinsics recorded for every source so comparisons between sources are valid.",
      "Courses are regenerated from committed seeds rather than stored. The seed list is "
      "the dataset."], bsize=16)
card(s, 10.2, 6.07, 8.9, 3.5, "Ethical",
     ["No camera frame is written to disk during flight, and no imagery leaves the "
      "operator's own equipment.",
      "Known failure modes are published in the model documentation rather than hidden.",
      "The headline figure is always estimated depth on held out courses. The perfect depth "
      "result is reported beside it as a stated ceiling."], dark=True, bsize=16)
tf = tbox(s, M, 9.85, 18.2, 0.5)
line(tf, "The latency budget is derived. At 2 m/s the aircraft travels 0.20 m during a "
         "100 ms delay and needs 0.67 m more to slow at 3 m/s squared, so stopping takes "
         "0.87 m before any margin. That is why FR3 sits above one metre.",
     14.5, MUTED, FL, first=True)

# =============================================================== 10. STAKEHOLDERS
s = slide("The least influence, the most at stake", "Stakeholder Management")
table(s, M, TOP, 12.2, 6.6,
      ["Stakeholder", "Role or interest", "Influence", "Main concern"],
      [["Bystanders on the ground", "Present beneath the aircraft. Agreed to nothing, gain nothing",
        "Very low", "Injury from a collision. Being recorded without consent"],
       ["Property owners", "Own the buildings and vehicles flown near",
        "Low in flight, high after", "Damage. Overflight without permission"],
       ["Operator and pilot", "Flies it and carries legal responsibility",
        "High", "Held responsible for software they did not write"],
       ["Regulators", "Set registration, certification and airspace rules",
        "Very high", "Unregistered operation. No demonstrated failsafe"],
       ["Institution", "Provides premises, carries liability",
        "High", "Injury on its premises. Reputational exposure"],
       ["Project team", "Design, build, evaluate, document",
        "High", "Delivering in a semester while reporting honestly"],
       ["Dataset providers", "Released the model and benchmarks under stated terms",
        "Low but binding", "Attribution. KITTI is non commercial"],
       ["Future maintainers", "Continue the work after we disperse",
        "Low now, high later", "Reproducing results without asking us"]],
      col_w=[2.6, 4.0, 1.9, 3.7], hsize=14, bsize=12.5)
card(s, 13.5, TOP, 5.6, 3.5, "Data subjects are not the affected community",
     ["Data subjects are the people in the benchmark photographs. Licence compliance "
      "protects them.",
      "The affected community is whoever stands beneath the aircraft. They gave nothing "
      "and face physical harm."], bsize=15)
card(s, 13.5, 6.07, 5.6, 3.5, "Conflicts we had to resolve",
     ["Team against bystanders. We gain from flying early, they carry the risk. So testing "
      "is staged.",
      "Team against assessors. A high number is rewarded, the honest one is lower. The "
      "reporting rule was fixed first."], dark=True, bsize=15)

# =============================================================== 11. FUTURE WORK
s = slide("The gap, the risk, the next moves", "Future Work and Next Steps")
picture(s, "figures/fig4_simulation.png", M, TOP + 0.02, 8.6)
card(s, M, 5.35, 8.6, 2.05, "Main justification and gap identified",
     ["Nobody has measured what avoidance performance costs when a distance sensor is "
      "replaced by a prediction, with everything else held constant. Our ablation does "
      "exactly that."], bsize=15)
card(s, M, 7.6, 8.6, 2.0, "Most important risk to address first",
     ["A reactive policy has no way out of a concave trap. With perfect depth we complete "
      "66 per cent of held out courses without contact, and most failures are timeouts, "
      "not crashes. That is the ceiling, so the 90 per cent target cannot be met as "
      "things stand."], dark=True, bsize=14.5)
card(s, 10.0, TOP, 9.1, 7.35, "Immediate next steps",
     ["1.  Put MiDaS into the loop and run the ablation. The policy already takes the depth "
      "map as an argument, so the planner needs no change.",
      "2.  Verify the scale and offset alignment before trusting any depth figure. Median "
      "scaling alone would make the model look broken when the alignment is at fault.",
      "3.  Measure whether the model runs on the Raspberry Pi. The architecture assumes it "
      "cannot, and that assumption has never been tested. If it can, the wireless link "
      "leaves the control loop entirely.",
      "4.  Resolve the 90 per cent target in writing. We prefer reporting a curve against "
      "obstacle density, which describes the method better than any single number.",
      "5.  Complete the mass and thrust budget before buying propulsion.",
      "6.  Confirm the regulatory position with the institution. This is a gate on outdoor "
      "flight, not paperwork to finish afterwards."], bsize=16, gap=10)

# =============================================================== 12. THANK YOU
s = prs.slides.add_slide(BLANK)
_bg(s)
rect(s, 1.65, 3.6, 2.25, 0.04, fill=INK)
tf = tbox(s, 1.65, 4.55, 16.8, 2.2)
line(tf, "Thank you", 76, INK, FS, first=True, space_after=14, spacing=1.0)
line(tf, "Questions welcome", 32, MUTED, FL)
rect(s, 1.65, 7.7, 4.5, 0.02, fill=RULE)
tf = tbox(s, 1.65, 8.1, 16.5, 1.4)
line(tf, "Siddharth Lama     Manit Aravind     Sahil Asopa     Riyaan Sheth",
     22, INK, FS, first=True, space_after=8)
line(tf, "Autonomous Drone Obstacle Avoidance Using Monocular Depth Estimation",
     16, MUTED, FL)
_n[0] += 1
num = tbox(s, W - 1.6, H - 0.85, 0.7, 0.35)
line(num, str(_n[0]), 13, MUTED, FL, first=True, align=PP_ALIGN.RIGHT)

OUT = "Project_Presentation_Unit1_2.pptx"
prs.save(OUT)
print("saved", OUT, "with", len(prs.slides), "slides")
