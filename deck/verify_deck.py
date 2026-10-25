"""Checks the deck without a renderer: bounds, estimated text overflow, fonts,
slide numbers, and a wireframe image so the layout can be eyeballed."""
import sys, math
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.util import Emu

F = "Project_Presentation_Unit1_2.pptx"
prs = Presentation(F)
W = prs.slide_width / 914400
H = prs.slide_height / 914400
inch = lambda v: (v or 0) / 914400
fails, warns = [], []


def est_height(tf, box_w):
    """Rough laid-out height in inches for a text frame of the given width."""
    total = 0.0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        pt = max([r.font.size.pt for r in p.runs if r.font.size] or [18])
        if not txt:
            total += pt * 1.2 / 72
            continue
        cpl = max(box_w * 144 / pt, 1)
        lines = math.ceil(len(txt) / cpl)
        ls = p.line_spacing or 1.0
        after = (p.space_after.pt if p.space_after else 0)
        total += lines * pt * 1.2 * ls / 72 + after / 72
    return total


def boxes_overlap(a, b, tol=0.05):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return (ax < bx + bw - tol and bx < ax + aw - tol
            and ay < by + bh - tol and by < ay + ah - tol)


for n, s in enumerate(prs.slides, 1):
    nums = []
    solid = []   # picture, table and card rectangles, which must not collide
    for sh in s.shapes:
        x, y, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
        if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
            fails.append(f"slide {n}: '{sh.shape_type}' out of bounds "
                         f"({x:.2f},{y:.2f}) {w:.2f}x{h:.2f}")
        if sh.shape_type == 13 or sh.has_table or (
                sh.has_text_frame and sh.text_frame.text.strip() and h > 1.0):
            solid.append(((x, y, w, h), str(sh.shape_type), sh))
        if sh.has_text_frame and sh.text_frame.text.strip():
            need = est_height(sh.text_frame, max(w - 0.85, 0.5))
            if need > h + 0.35 and h > 0.6:
                warns.append(f"slide {n}: text may overflow, needs ~{need:.2f} in "
                             f"of {h:.2f} in  [{sh.text_frame.text[:38]!r}]")
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name and not r.font.name.startswith("Inter"):
                        fails.append(f"slide {n}: font {r.font.name}")
                    if r.text.strip().isdigit() and len(r.text.strip()) <= 2:
                        nums.append(r.text.strip())
    for i in range(len(solid)):
        for j in range(i + 1, len(solid)):
            if boxes_overlap(solid[i][0], solid[j][0]):
                fails.append(f"slide {n}: {solid[i][1]} overlaps {solid[j][1]} "
                             f"at {tuple(round(q,2) for q in solid[i][0])} and "
                             f"{tuple(round(q,2) for q in solid[j][0])}")
    if str(n) not in nums:
        fails.append(f"slide {n}: no page number")

print(f"slides {len(prs.slides)} | size {W:.2f} x {H:.2f} in")
print(f"tables {sum(1 for s in prs.slides for sh in s.shapes if sh.has_table)} | "
      f"pictures {sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)}")
print()
if warns:
    print("WARNINGS")
    for w_ in warns:
        print("  -", w_)
    print()
if fails:
    print("FAILURES")
    for f_ in fails:
        print("  -", f_)
    sys.exit(1)
print("structural checks passed")

# wireframe so the layout can actually be seen
cols = 3
rows = math.ceil(len(prs.slides) / cols)
fig, axes = plt.subplots(rows, cols, figsize=(cols * 4.4, rows * 2.6))
for i, (s, ax) in enumerate(zip(prs.slides, axes.ravel()), 1):
    ax.add_patch(Rectangle((0, 0), W, H, fc="#EFEEEC", ec="none"))
    for sh in s.shapes:
        x, y, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
        if sh.shape_type == 13:
            fc, ec = "#cfd6dd", "#6b7480"
        elif sh.has_table:
            fc, ec = "#e3ded8", "#403E3E"
        elif sh.has_text_frame and sh.text_frame.text.strip():
            fc, ec = "#ffffff", "#8A8684"
        else:
            fc, ec = "#403E3E", "none"
        ax.add_patch(Rectangle((x, H - y - h), w, h, fc=fc, ec=ec, lw=0.5, alpha=0.9))
        if sh.has_text_frame and sh.text_frame.text.strip():
            ax.text(x + 0.15, H - y - 0.45, sh.text_frame.text.strip()[:30],
                    fontsize=3.4, color="#403E3E")
    ax.set_xlim(0, W); ax.set_ylim(0, H); ax.set_aspect("equal")
    ax.set_title(f"slide {i}", fontsize=7); ax.axis("off")
for ax in axes.ravel()[len(prs.slides):]:
    ax.axis("off")
fig.tight_layout()
fig.savefig("wireframe.png", dpi=150)
print("wrote wireframe.png")
