"""Structural checks. There is no Word renderer on this machine, so verify the
XML rather than trusting the appearance."""
import re, sys, zipfile
from docx import Document
from docx.shared import RGBColor

F = "Project_Report_Unit1_2_Monocular_Depth_VTOL.docx"
doc = Document(F)
fails = []

def collect_text():
    out = []
    for p in doc.paragraphs:
        out.append(p.text)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    out.append(p.text)
    return out

texts = collect_text()
blob = "\n".join(texts)

# 1. forbidden characters
for name, ch in [("em dash", "—"), ("en dash", "–"), ("semicolon", ";"),
                 ("smart quote", "“"), ("smart apostrophe", "’")]:
    hits = [t for t in texts if ch in t]
    if hits:
        fails.append(f"{name} found in {len(hits)} block(s): {hits[0][:90]!r}")

# 2. fonts and colour on every run
bad_font, bad_colour, sup_sub = set(), set(), 0
def check(runs, where):
    global sup_sub
    for r in runs:
        if r.text.strip():
            # Courier New is deliberate, for command and formula blocks
            if r.font.name not in ("Times New Roman", "Courier New"):
                bad_font.add((where, r.font.name))
            c = r.font.color
            if c is None or c.rgb != RGBColor(0, 0, 0):
                bad_colour.add((where, str(c.rgb) if c and c.rgb else "unset"))
            if r.font.superscript or r.font.subscript:
                sup_sub += 1
for p in doc.paragraphs:
    check(p.runs, "body")
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                check(p.runs, "table")
if bad_font: fails.append(f"unexpected fonts: {sorted(bad_font)[:4]}")
if bad_colour: fails.append(f"non black runs: {sorted(bad_colour)[:4]}")
if sup_sub: fails.append(f"{sup_sub} superscript or subscript runs")

# 3. no headers or footers with content
with zipfile.ZipFile(F) as z:
    hf = [n for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml", n)]
    for n in hf:
        body = z.read(n).decode("utf8", "ignore")
        if re.search(r"<w:t[^>]*>[^<]+</w:t>", body):
            fails.append(f"{n} contains visible text")

# 4. headings are bold
heads = [p for p in doc.paragraphs
         if p.runs and (p.text.startswith("CHAPTER") or re.match(r"^\d+\.\d+ ", p.text)
                        or p.text in ("DEDICATION","ABSTRACT","ACKNOWLEDGEMENTS","PREFACE",
                                      "INDEX","LIST OF FIGURES","LIST OF TABLES"))]
notbold = [p.text[:50] for p in heads if not all(r.font.bold for r in p.runs if r.text.strip())]
if notbold: fails.append(f"headings not bold: {notbold}")

# 5. required structure present
need = ["DEDICATION","ABSTRACT","ACKNOWLEDGEMENTS","PREFACE","INDEX","LIST OF FIGURES",
        "LIST OF TABLES","CHAPTER I.","CHAPTER II.","CHAPTER III.","CHAPTER IV.",
        "CHAPTER V.","CHAPTER VI.","CHAPTER VII.","CHAPTER VIII.","CHAPTER IX.",
        "CHAPTER X.","CHAPTER XI."]
for n in need:
    if not any(t.startswith(n) for t in texts):
        fails.append(f"missing section: {n}")

# 6. everyone named
for name in ["Siddharth Lama","Manit Aravind","Sahil Asopa","Riyaan Sheth"]:
    if name not in blob: fails.append(f"student missing: {name}")

# 7. images
with zipfile.ZipFile(F) as z:
    imgs = [n for n in z.namelist() if n.startswith("word/media/")]
words = len(blob.split())
print(f"tables {len(doc.tables)} | images {len(imgs)} | headings {len(heads)} | words {words}")
print(f"estimated pages at ~380 words plus tables and figures: {round(words/380 + len(doc.tables)*0.35 + len(imgs)*0.5)}")
print()
if fails:
    print("FAILED CHECKS")
    for f in fails: print("  -", f)
    sys.exit(1)
print("ALL CHECKS PASSED")
